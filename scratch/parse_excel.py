import openpyxl
import json

wb = openpyxl.load_workbook('India_Top_500_Tourist_Destinations.xlsx')
print("Sheet names:", wb.sheetnames)

sheet = wb.active
rows = list(sheet.iter_rows(values_only=True))

if rows:
    headers = rows[0]
    print("Headers:", headers)
    print("Total rows:", len(rows) - 1)
    print("\nSample Rows:")
    for row in rows[1:6]:
        print(row)
